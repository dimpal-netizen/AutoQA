"""Drafting bug reports. No network.

A bug report is read by someone who will not open AutoQA, so the guards here
are about what reaches a ticket: a title that fits on one line, reproduction
steps that are actually present, and a severity that agrees with the analysis
rather than contradicting it in the same panel.
"""

import pytest

from app.ai.schemas import DraftedBug
from app.models.enums import BugStatus, Severity
from app.reports.bugs import COLUMNS
from app.services.bug_service import _line, _text

#: Column number by name, so a test says what it reads rather than counting.
COLUMN = {name: i for i, name in enumerate(COLUMNS, start=1)}


def drafted(**overrides) -> DraftedBug:
    base = dict(
        title="Login rejects an email with trailing spaces",
        description="Signing in fails when the email has whitespace around it.",
        steps_to_reproduce=[
            "Open https://example.test/login",
            "Type ' user@example.com ' into the Email field",
            "Type the correct password",
            "Click Login",
        ],
        expected="The user is signed in and taken to the dashboard.",
        actual="The page stays on /login and shows 'Invalid email or password'.",
        severity="high",
        priority="high",
    )
    base.update(overrides)
    return DraftedBug(**base)


# ---------------------------------------------------------------------------
# Text going into a ticket
# ---------------------------------------------------------------------------
def test_a_title_is_forced_onto_one_line():
    """Titles land in a backlog row; a newline there breaks the whole list."""
    messy = _line("Login\n\n  fails   badly\t\tsometimes", 255)

    assert messy == "Login fails badly sometimes"
    assert "\n" not in messy


def test_a_title_is_truncated_to_fit_the_column():
    assert len(_line("x" * 900, 255)) == 255


def test_control_characters_are_stripped():
    assert _line("Login\x00 fails\x07", 255) == "Login fails"


def test_body_text_keeps_its_paragraphs():
    """Unlike the title, a description is allowed to breathe."""
    body = _text("First line.\n\nSecond line.")

    assert "\n\n" in body


@pytest.mark.parametrize("empty", ["", "   ", None])
def test_empty_text_does_not_become_the_string_none(empty):
    assert _text(empty) == ""
    assert _line(empty, 255) == ""


# ---------------------------------------------------------------------------
# The shape of what the model returns
# ---------------------------------------------------------------------------
def test_a_draft_carries_everything_a_developer_needs():
    bug = drafted()

    assert bug.title and bug.description
    assert bug.steps_to_reproduce, "a bug nobody can reproduce gets closed"
    assert bug.expected and bug.actual


def test_reproduction_steps_are_ordered_and_concrete():
    """The steps decide whether a bug is fixed or closed as 'cannot reproduce'."""
    steps = drafted().steps_to_reproduce

    assert steps[0].startswith("Open http")
    assert any("user@example.com" in step for step in steps), (
        "steps must carry the actual values used, not placeholders"
    )


def test_a_drafted_report_starts_as_a_draft():
    """An AI-written ticket nobody reviewed must not look raised."""
    assert BugStatus.DRAFT.value == "draft"
    assert BugStatus.DRAFT is not BugStatus.OPEN


def test_severity_shares_one_vocabulary_with_analysis():
    """The bug and the analysis sit next to each other; two scales would
    contradict each other on the same screen."""
    from app.ai.analyser import severity_of

    assert severity_of("critical") is Severity.CRITICAL
    assert severity_of(drafted().severity) is Severity.HIGH


# ---------------------------------------------------------------------------
# Every bug in one workbook
#
# The per-failure report is a ticket: one problem, handed to one developer.
# A project has dozens, and the questions asked of them are different ones —
# what is still open, what is critical, how much of this is a single browser.
# None of those can be answered by opening tickets one at a time, which was the
# only way to read them.
# ---------------------------------------------------------------------------
def bug(**overrides):
    """One register row. The workbook takes these however they were built."""
    from datetime import datetime

    from app.reports.bugs import BugRow

    base = dict(
        title="Login rejects an email with trailing spaces",
        description="Signing in fails when the email has whitespace around it.",
        steps_to_reproduce=["Open /login", "Type ' a@b.com '", "Click Login"],
        expected="The user is signed in.",
        actual="The page stays on /login.",
        case_name="Login with a valid account",
        browsers="chromium",
        environment={"url": "https://x.test/login"},
        severity=Severity.HIGH,
        priority=Severity.HIGH,
        status=BugStatus.OPEN,
        reported_on=datetime(2026, 8, 6, 11, 0),
        drafted=True,
    )
    base.update(overrides)
    return BugRow(**base)


def workbook(bugs, project_name: str = "Homeske"):
    """Build and read back, so a test asserts on the file rather than the code."""
    import io

    from openpyxl import load_workbook

    from app.reports.bugs import build_bug_workbook

    raw = build_bug_workbook(bugs, project_name=project_name)
    return load_workbook(io.BytesIO(raw)).active


def test_every_bug_gets_a_row():
    sheet = workbook([bug(), bug(), bug()])

    # Eight rows of title and header block, then one per bug.
    assert sheet.max_row == 8 + 3


def test_the_worst_bug_is_the_first_one_you_read():
    """A register sorted by id makes you read all of it to find the argument."""
    sheet = workbook([
        bug(title="Low one", severity=Severity.LOW),
        bug(title="Critical one", severity=Severity.CRITICAL),
        bug(title="Medium one", severity=Severity.MEDIUM),
    ])

    titles = [
        sheet.cell(row=9 + i, column=COLUMN["Bug Summary"]).value for i in range(3)
    ]
    assert titles == ["Critical one", "Medium one", "Low one"]


def test_the_steps_are_numbered_one_per_line():
    """A bug nobody can reproduce gets closed, so this column decides its fate."""
    sheet = workbook([bug()])
    steps = sheet.cell(row=9, column=COLUMN["Steps to Reproduce"]).value

    assert steps.splitlines() == ["1. Open /login", "2. Type ' a@b.com '", "3. Click Login"]


def test_the_description_line_says_how_many_there_are():
    """Their template has one Description line rather than a block of counts,
    so the number goes there instead of being dropped."""
    sheet = workbook([bug(), bug(), bug()])
    header = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
              for r in range(2, 7)}

    assert "3 bug(s)" in header["Description"]


def test_a_bug_whose_run_was_deleted_still_exports():
    """`result_id` is nulled when a run goes. The steps were copied in already."""
    sheet = workbook([bug(case_name="")])

    # openpyxl reads an empty cell back as None; either way the column is blank
    # rather than the export having failed.
    assert not sheet.cell(row=9, column=COLUMN["Test case ID"]).value
    assert sheet.cell(row=9, column=COLUMN["Steps to Reproduce"]).value


def test_the_sheet_filters_and_freezes_so_it_can_be_triaged():
    """"Show me open criticals" is the question this file exists to answer."""
    sheet = workbook([bug() for _ in range(5)])

    from openpyxl.utils import get_column_letter

    assert sheet.freeze_panes == "A9"
    # Derived, so adding a column moves the filter instead of breaking it.
    last = get_column_letter(len(COLUMNS))
    assert sheet.auto_filter.ref == f"A8:{last}13"


def service_with(failures, analyses=None, steps=None):
    """A BugService with only the parts `_rows_from_failures` touches."""
    from app.services.bug_service import BugService

    service = BugService.__new__(BugService)
    found = analyses or {}
    service.analyses = type("A", (), {"latest_for_result": lambda _s, rid: found.get(rid)})()
    service._steps_for = lambda _result: steps or []
    return service


class FakeStep:
    def __init__(self, description: str) -> None:
        self.description = description
        self.input_data = None


class FakeFailure:
    def __init__(self, case_id: int, browser: str, **overrides) -> None:
        from datetime import datetime

        from app.models.enums import Browser, ResultStatus

        self.id = case_id * 10 + len(browser)
        self.test_case_id = case_id
        self.browser = Browser(browser)
        self.status = ResultStatus.FAILED
        self.case_name = overrides.get("case_name", "Successful registration")
        self.error_message = overrides.get("error_message", "AssertionError: hidden")
        self.failed_step = overrides.get("failed_step", 7)
        self.created_at = datetime(2026, 8, 6, 9, 0)


class FakeAnalysis:
    def __init__(self, **overrides) -> None:
        base = dict(
            expected="The one-time-password popup should have opened.",
            actual="The form stayed put with a red message under Mobile Number.",
            root_cause="The phone number is rejected for the selected country.",
            severity=Severity.CRITICAL,
            priority=Severity.HIGH,
        )
        base.update(overrides)
        for key, value in base.items():
            setattr(self, key, value)


def service_with(failures, analyses=None, steps=None):
    """A BugService with only the parts `_rows_from_failures` touches."""
    from app.services.bug_service import BugService

    service = BugService.__new__(BugService)
    found = analyses or {}
    service.analyses = type("A", (), {"latest_for_result": lambda _s, rid: found.get(rid)})()
    service._steps_for = lambda _result: steps or []
    return service


PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c6360000002000100fdff03fa0000000049454e44ae426082"
)

def test_a_failure_nobody_drafted_becomes_a_row():
    service = service_with([])
    rows = service._rows_from_failures([FakeFailure(1, "chromium")])

    assert len(rows) == 1
    assert rows[0].drafted is False
    assert rows[0].case_name == "Successful registration"


def test_one_bug_not_one_per_browser():
    """Red in two browsers is one bug affecting two, not two bugs."""
    service = service_with([])
    rows = service._rows_from_failures([
        FakeFailure(1, "chromium"), FakeFailure(1, "firefox")
    ])

    assert len(rows) == 1
    assert rows[0].browsers == "chromium, firefox"


def test_the_analysis_writes_the_row_when_there_is_one():
    failure = FakeFailure(1, "chromium")
    service = service_with([], analyses={failure.id: FakeAnalysis()})
    row = service._rows_from_failures([failure])[0]

    assert row.title == "The phone number is rejected for the selected country."
    assert row.expected.startswith("The one-time-password popup")
    assert row.severity is Severity.CRITICAL


def test_without_an_analysis_the_row_still_says_something_useful():
    """No AI has run, and the export must not wait for one."""
    service = service_with([], steps=[FakeStep("Open the form"), FakeStep("Submit")])
    row = service._rows_from_failures([FakeFailure(1, "chromium")])[0]

    assert "Successful registration" in row.title
    assert row.actual == "AssertionError: hidden"
    assert row.steps_to_reproduce == ["Open the form", "Submit"]
    assert row.severity is Severity.MEDIUM


def document(bugs, project_name: str = "Homeske"):
    """Build and read back, so a test asserts on the file rather than the code."""
    import io

    from docx import Document

    from app.reports.bugs_doc import build_bug_document

    return Document(io.BytesIO(build_bug_document(bugs, project_name=project_name)))


def text_of(doc) -> str:
    """Every word in the document, headings and table cells alike.

    The tables matter as much as the paragraphs now: an issue is laid out the
    way Mantis lays one out, a label column beside a value column, so reading
    only `doc.paragraphs` finds the headings and none of the content.
    """
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_the_screenshot_is_in_the_document():
    """The one thing a spreadsheet cell cannot hold."""
    doc = document([bug(screenshot=PNG)])

    assert len(doc.inline_shapes) == 1


def test_every_bug_brings_its_own_screenshot():
    doc = document([bug(screenshot=PNG), bug(screenshot=PNG), bug(screenshot=PNG)])

    assert len(doc.inline_shapes) == 3


def test_a_failure_with_no_screenshot_says_so_rather_than_going_quiet():
    """A missing picture is worth a sentence; a blank space reads as a bug."""
    doc = document([bug(screenshot=None)])

    assert "No screenshot was captured" in text_of(doc)


def test_an_unreadable_image_costs_one_picture_not_the_report():
    """Truncated on disk, or a format Word declines. The other bugs still ship."""
    doc = document([bug(title="First", screenshot=b"not an image"), bug(title="Second", screenshot=PNG)])

    assert "could not be embedded" in text_of(doc)
    assert "Second" in text_of(doc)
    assert len(doc.inline_shapes) == 1


def test_the_worst_bug_is_the_first_one_you_read_here_too():
    doc = document([
        bug(title="Low one", severity=Severity.LOW),
        bug(title="Critical one", severity=Severity.CRITICAL),
    ])
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]

    assert headings[0].endswith("Critical one")


def test_the_prose_a_developer_needs_is_all_there():
    doc = document([bug()])
    body = text_of(doc)

    assert "Login rejects an email with trailing spaces" in body
    assert "The user is signed in." in body        # expected
    assert "The page stays on /login." in body     # actual
    assert "Open /login" in body                   # steps to reproduce


def test_an_empty_field_does_not_leave_a_dangling_label():
    """"Expected:" with nothing after it reads as a tool that failed to fill in."""
    doc = document([bug(expected="", case_name="")])
    body = text_of(doc)

    assert "Expected:" not in body
    assert "Test case:" not in body


def test_an_undrafted_bug_says_nobody_has_reviewed_it():
    """Mantis has no field for it, so it is said in Additional Information.

    Left out entirely, an unreviewed row would read as a triaged one - and a
    register that presents AutoQA's guesses as somebody's findings is how a
    team learns to distrust the register.
    """
    doc = document([bug(drafted=False)])

    assert "not yet reviewed" in text_of(doc)


def test_the_summary_counts_what_nobody_has_looked_at():
    doc = document([bug(drafted=True), bug(drafted=False)])

    assert "2 bug(s)." in text_of(doc)
    assert "1 have not been reviewed" in text_of(doc)


def test_both_exports_are_built_from_the_same_row():
    """One description of a bug, so the two files cannot drift apart."""
    from app.reports.bugs import BugRow
    from app.reports.bugs_doc import build_bug_document

    one = bug()
    assert isinstance(one, BugRow)
    assert build_bug_document([one], project_name="P")
    assert workbook([one])

# ---------------------------------------------------------------------------
# Two files, two formats, on purpose
#
# The workbook is the team's own bug-report template — the sheet they already
# circulate. The document is Mantis's Report Issue form, because that is what
# it is for: filling in the tracker. A register that arrives in a shape nobody
# recognises gets re-typed into one that is, so matching each audience beats
# making both files match each other.
# ---------------------------------------------------------------------------
def test_the_sheet_is_the_teams_own_template():
    """Their columns, their order, their spelling of "Serverity" — corrected,
    it would stop matching the template somebody pastes into."""
    assert COLUMNS == [
        "Bug ID",
        "Bug Summary",
        "Serverity",
        "Steps to Reproduce",
        "Expected Result",
        "Actual Results",
        "Browser/OS used",
        "Status",
        "Test case ID",
        "Comments/Screen shots",
    ]


def test_the_sheet_opens_with_the_five_lines_their_template_opens_with():
    sheet = workbook([bug()], project_name="Homeske Web Application")
    labels = [sheet.cell(row=r, column=1).value for r in range(2, 7)]

    assert labels == [
        "Project Name",
        "Module Name",
        "Description",
        "Bug  Reported by",
        "Reported Date",
    ]
    assert sheet.cell(row=2, column=2).value == "Homeske Web Application"


def test_the_module_is_left_for_a_person_to_fill_in():
    """A project has several and nothing in a run says which one a failure
    belongs to. Blank is a box somebody fills; a guess is one they check."""
    sheet = workbook([bug()])

    assert not sheet.cell(row=3, column=2).value


def test_severity_uses_the_words_on_their_sheet():
    """High / Medium / Low / Blocker, which is what the template's own hint
    cell lists. "critical" is not one of them."""
    sheet = workbook([bug(severity=Severity.CRITICAL)])

    assert sheet.cell(row=9, column=COLUMN["Serverity"]).value == "Blocker"


def test_browser_and_os_share_one_cell_because_the_column_asks_for_both():
    sheet = workbook([bug(environment={"os": "Windows 11"})])

    assert sheet.cell(row=9, column=COLUMN["Browser/OS used"]).value == (
        "chromium / Windows 11"
    )


def test_the_comments_column_says_where_the_screenshot_is():
    """Their example is a URL. AutoQA has the picture rather than a link to
    one, so this points at the file that holds it instead of sitting blank as
    though there were nothing."""
    sheet = workbook([bug(screenshot=b"x")])

    assert "Word report" in sheet.cell(row=9, column=COLUMN["Comments/Screen shots"]).value


def test_an_unreviewed_row_says_so_in_the_comments():
    """Mantis has no field for it and neither does their template, so it goes
    where a person would write it. Left out, an unreviewed row reads as a
    triaged one."""
    service = service_with([])
    sheet = workbook(service._rows_from_failures([FakeFailure(1, "chromium")]))

    assert "not yet reviewed" in (
        sheet.cell(row=9, column=COLUMN["Comments/Screen shots"]).value or ""
    )


# --- the document is still Mantis ------------------------------------------
def test_the_document_reads_the_way_mantis_shows_an_issue():
    """Identity first, then the classification, then the boxes somebody types
    into — the order the tracker itself puts them in."""
    labels = [row.cells[0].text for row in document([bug()]).tables[0].rows]

    assert labels[:6] == [
        "ID", "Project", "Category", "View Status", "Date Submitted", "Last Update",
    ]
    assert labels[-1] == "Tags"


def _box(doc, label: str) -> str:
    row = next(r for r in doc.tables[0].rows if r.cells[0].text == label)
    return "\n".join(p.text for p in row.cells[1].paragraphs)


def test_the_document_translates_severity_into_mantis_words():
    """AutoQA says critical; Mantis has no such severity. Handing over a word
    the tracker does not know means somebody picks one by hand, per issue."""
    doc = document([bug(severity=Severity.CRITICAL, priority=Severity.CRITICAL)])

    assert _box(doc, "Severity").strip() == "block"
    assert _box(doc, "Priority").strip() == "immediate"


def test_a_drafted_bug_is_new_rather_than_acknowledged():
    """Mantis's `acknowledged` and `confirmed` both claim a person has looked.
    Nobody has: AutoQA wrote it."""
    doc = document([bug(status=BugStatus.DRAFT)])

    assert _box(doc, "Status").strip() == "new"
    assert _box(doc, "Resolution").strip() == "open"


def test_the_document_keeps_the_headings_mantis_puts_in_its_boxes():
    """The form arrives holding Summary / Expected Result / Actual Result, and
    Browser Used / Device used. A failure knows them all, so they arrive
    completed — and the headings stay, so the box still looks like the one the
    team fills in by hand."""
    doc = document([bug()])

    description = _box(doc, "Description")
    assert "Summary:" in description
    assert "Expected Result: The user is signed in." in description
    assert "Actual Result: The page stays on /login." in description
    assert "Browser Used: chromium" in _box(doc, "Additional Information")


def test_the_boxes_only_a_person_can_fill_are_present_and_empty():
    """A form with a row missing is harder to work through than one with a row
    to skip."""
    doc = document([bug()])

    assert _box(doc, "Select Profile").strip() == ""
    assert _box(doc, "Assigned To").strip() == ""
