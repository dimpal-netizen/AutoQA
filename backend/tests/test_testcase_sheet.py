"""Exporting a suite as a QA test-case workbook. No database, no network.

The property that matters is that the sheet is usable by a person who has never
seen AutoQA: every column a QA reviewer expects is present, in order, and the
ones AutoQA can answer are answered rather than left for someone to retype.

Read back through openpyxl rather than asserted against bytes — the test should
see what Excel will see.
"""

import io
from dataclasses import dataclass, field
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from app.models.enums import ActionType, CaseCategory, CasePriority, ResultStatus
from app.reports.testcases import COLUMNS, build_testcase_workbook


@dataclass
class FakeStep:
    description: str
    action: ActionType = ActionType.CLICK
    input_data: str | None = None
    expected_result: str | None = None


@dataclass
class FakeCase:
    name: str
    function_name: str
    category: CaseCategory = CaseCategory.RECORDED
    priority: CasePriority = CasePriority.HIGH
    steps: list = field(default_factory=list)


@dataclass
class FakeProject:
    name: str = "Homeske"


@dataclass
class FakeSuite:
    id: int = 1
    name: str = "Subscription Payment"
    description: str | None = "M-Pesa subscription payment"
    created_at: datetime = datetime(2026, 6, 5, tzinfo=UTC)
    project: FakeProject = field(default_factory=FakeProject)
    cases: list = field(default_factory=list)


@dataclass
class FakeRun:
    id: int = 9
    finished_at: datetime | None = datetime(2026, 7, 30, tzinfo=UTC)


@dataclass
class FakeResult:
    function_name: str
    status: ResultStatus = ResultStatus.PASSED
    error_message: str | None = None


@pytest.fixture
def cases() -> list[FakeCase]:
    return [
        FakeCase(
            name="Verify successful M-Pesa payment for Monthly plan",
            function_name="test_monthly_payment",
            category=CaseCategory.RECORDED,
            priority=CasePriority.HIGH,
            steps=[
                FakeStep("Open the subscribe page", ActionType.NAVIGATE, "/subscribe"),
                FakeStep("Enter M-Pesa number", ActionType.INPUT, "+254700111222"),
                FakeStep(
                    "Submit",
                    ActionType.CLICK,
                    expected_result="Subscription should be activated.",
                ),
            ],
        ),
        FakeCase(
            name="Verify payment fails with an invalid M-Pesa number",
            function_name="test_invalid_number",
            category=CaseCategory.NEGATIVE,
            priority=CasePriority.MEDIUM,
            steps=[FakeStep("Enter invalid number", ActionType.INPUT, "12345abc")],
        ),
        FakeCase(
            name="SQL payload in the number field is rejected",
            function_name="test_sql_payload",
            category=CaseCategory.SECURITY,
            priority=CasePriority.CRITICAL,
        ),
    ]


def worksheet(cases, **kwargs):
    """Build the workbook and hand back the sheet Excel would open."""
    suite = FakeSuite(cases=cases)
    payload = build_testcase_workbook(suite, cases, project_name="Homeske", **kwargs)
    return load_workbook(io.BytesIO(payload)).active


def sheet(cases, **kwargs) -> list[list[str]]:
    """Every cell as text, so the existing assertions read unchanged."""
    return [
        ["" if cell is None else str(cell) for cell in row]
        for row in worksheet(cases, **kwargs).iter_rows(values_only=True)
    ]


def body(table: list[list[str]]) -> list[list[str]]:
    """Everything below the column header row."""
    header = table.index(COLUMNS)
    return table[header + 1 :]


def test_header_block_names_the_project_and_module(cases):
    table = sheet(cases)
    flat = {row[0]: row[1] for row in table if len(row) > 1 and row[0]}

    assert flat["Project Name"] == "Homeske"
    assert flat["Module Name"] == "Subscription Payment"
    assert flat["Creation Date"] == "05-06-2026"


def test_every_expected_column_is_present_and_in_order(cases):
    assert COLUMNS in sheet(cases)


def test_case_ids_are_prefixed_and_sequential(cases):
    ids = [row[0] for row in body(sheet(cases))]
    # Homeske + Subscription Payment -> HOM-SP.
    assert ids == ["HOM-SP-001", "HOM-SP-002", "HOM-SP-003"]


def test_positive_and_negative_are_derived_from_the_category(cases):
    column = COLUMNS.index("Positive/Negative")
    assert [row[column] for row in body(sheet(cases))] == [
        "Positive",  # recorded
        "Negative",  # negative
        "Negative",  # security
    ]


def test_security_cases_are_not_labelled_functional(cases):
    column = COLUMNS.index("Type")
    assert [row[column] for row in body(sheet(cases))] == [
        "Functional",
        "Functional",
        "Security",
    ]


def test_steps_are_numbered_one_per_line(cases):
    column = COLUMNS.index("Test Steps")
    assert body(sheet(cases))[0][column] == (
        "1. Open the subscribe page\n2. Enter M-Pesa number\n3. Submit"
    )


def test_input_data_names_the_field_it_belongs_to(cases):
    """A bare value is useless to whoever re-tests this by hand."""
    column = COLUMNS.index("Input Data")
    assert body(sheet(cases))[0][column] == (
        "Open the subscribe page: /subscribe\nEnter M-Pesa number: +254700111222"
    )


def test_a_case_with_no_input_says_so_rather_than_being_blank(cases):
    column = COLUMNS.index("Input Data")
    assert body(sheet(cases))[2][column] == "N/A"


def test_execution_columns_are_blank_without_a_run(cases):
    table = body(sheet(cases))
    for name in ("Actual Results", "Status", "Test Execution Date"):
        column = COLUMNS.index(name)
        assert all(row[column] == "" for row in table)


def test_a_run_fills_in_the_execution_columns(cases):
    table = body(
        sheet(
            cases,
            run=FakeRun(),
            results=[
                FakeResult("test_monthly_payment"),
                FakeResult(
                    "test_invalid_number",
                    ResultStatus.FAILED,
                    "TimeoutError: Locator.click timed out",
                ),
            ],
        )
    )

    status = COLUMNS.index("Status")
    actual = COLUMNS.index("Actual Results")
    date = COLUMNS.index("Test Execution Date")

    assert table[0][status] == "Pass"
    assert table[0][actual] == "As expected."
    assert table[0][date] == "30-07-2026"

    assert table[1][status] == "Fail"
    assert "TimeoutError" in table[1][actual]

    # The third case did not run, so its columns stay empty rather than
    # claiming a result nobody produced.
    assert table[2][status] == ""
    assert table[2][date] == ""


def test_a_failure_on_any_browser_beats_a_pass(cases):
    """One Status column, three browsers. Passing two out of three is not a pass."""
    table = body(
        sheet(
            cases,
            run=FakeRun(),
            results=[
                FakeResult("test_monthly_payment", ResultStatus.PASSED),
                FakeResult("test_monthly_payment", ResultStatus.FAILED, "boom"),
                FakeResult("test_monthly_payment", ResultStatus.PASSED),
            ],
        )
    )
    assert table[0][COLUMNS.index("Status")] == "Fail"


def test_columns_a_human_owns_are_left_empty(cases):
    """Retesting and peer review are judgements, not data we can invent."""
    table = body(sheet(cases, run=FakeRun(), results=[]))
    for name in ("Retesting Status", "Retesting Date", "Peer Review", "Comments"):
        column = COLUMNS.index(name)
        assert all(row[column] == "" for row in table)


def test_every_row_has_a_cell_for_every_column(cases):
    """A short row silently shifts every value after it into the wrong column."""
    for row in body(sheet(cases)):
        assert len(row) == len(COLUMNS)


def test_a_suite_with_no_cases_still_produces_a_usable_sheet():
    table = sheet([])
    assert COLUMNS in table
    assert body(table) == []


# ---------------------------------------------------------------------------
# The formatting, which is half of why it is a workbook and not a CSV
# ---------------------------------------------------------------------------
def test_it_is_a_real_workbook(cases):
    """Not a CSV with a different extension — Excel refuses those with a warning."""
    payload = build_testcase_workbook(
        FakeSuite(cases=cases), cases, project_name="Homeske"
    )
    assert payload[:2] == b"PK"  # xlsx is a zip


def test_the_header_row_is_frozen_and_filterable(cases):
    """A hundred-row sheet is unusable without both."""
    ws = worksheet(cases)
    assert ws.freeze_panes == "A9"
    assert ws.auto_filter.ref.startswith("A8:")


def test_positive_and_negative_cases_are_tinted_apart(cases):
    """The split a reviewer opens the sheet for, visible without reading."""
    ws = worksheet(cases)
    first_data_row = 9

    positive = ws.cell(row=first_data_row, column=1).fill.fgColor.rgb
    negative = ws.cell(row=first_data_row + 1, column=1).fill.fgColor.rgb
    assert positive != negative


def test_long_text_wraps_instead_of_running_off(cases):
    """Steps are the widest column and the reason CSV was unreadable."""
    ws = worksheet(cases)
    steps = ws.cell(row=9, column=COLUMNS.index("Test Steps") + 1)
    assert steps.alignment.wrap_text is True
    assert steps.alignment.vertical == "top"


def test_every_column_has_a_width_set(cases):
    ws = worksheet(cases)
    from openpyxl.utils import get_column_letter

    for index in range(1, len(COLUMNS) + 1):
        assert ws.column_dimensions[get_column_letter(index)].width
